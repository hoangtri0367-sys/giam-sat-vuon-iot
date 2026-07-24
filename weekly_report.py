"""
weekly_report.py
----------------
Tổng hợp dữ liệu phân tích ảnh cả tuần → xuất file .xlsx
và gửi qua email (SMTP).

Cách dùng:
  1. Tự động: APScheduler gọi generate_and_send_weekly_report() mỗi Chủ nhật 20:00
  2. Thủ công: POST /api/export-weekly  (xem flask_routes.py)

Cấu hình email → file .env hoặc biến môi trường Railway:
  EMAIL_SENDER, EMAIL_PASSWORD, EMAIL_RECIPIENT, SMTP_HOST, SMTP_PORT
"""

import os
import io
import json
import smtplib
import logging
from datetime import datetime, timedelta
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ────────────────────────────────────────────────
# CONFIG (đọc từ env)
# ────────────────────────────────────────────────

EMAIL_SENDER    = os.environ.get("EMAIL_SENDER",    "your_email@gmail.com")
EMAIL_PASSWORD  = os.environ.get("EMAIL_PASSWORD",  "your_app_password")
EMAIL_RECIPIENT = os.environ.get("EMAIL_RECIPIENT", "recipient@gmail.com")
SMTP_HOST       = os.environ.get("SMTP_HOST",       "smtp.gmail.com")
SMTP_PORT       = int(os.environ.get("SMTP_PORT",   "587"))

# ────────────────────────────────────────────────
# STYLE HELPERS
# ────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _font(bold=False, size=10, color="000000") -> Font:
    return Font(bold=bold, size=size, color=color, name="Calibri")

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

HEADER_FILL   = _fill("2E7D32")   # xanh đậm
SUBHEAD_FILL  = _fill("A5D6A7")   # xanh nhạt
ALT_ROW_FILL  = _fill("F1F8E9")   # xanh rất nhạt
WARNING_FILL  = _fill("FFF9C4")   # vàng
DANGER_FILL   = _fill("FFCDD2")   # đỏ nhạt
HEALTHY_FILL  = _fill("C8E6C9")   # xanh lành mạnh


# ────────────────────────────────────────────────
# HÀM CHÍNH
# ────────────────────────────────────────────────

def generate_weekly_excel(records: list, week_start: datetime = None, sensor_records: list = None) -> bytes:
    """
    Tạo file Excel từ danh sách records phân tích ảnh trong tuần.

    Parameters
    ----------
    records : list of dict  – mỗi phần tử là kết quả của analyze_image()
                               đã được lưu vào DB/file, có thêm key "timestamp"
    week_start : datetime   – ngày đầu tuần (None = tự tính 7 ngày trước)
    sensor_records : list of dict – dữ liệu cảm biến môi trường (bảng sensor_data:
                               timestamp, temperature, humidity, soil_moist, light_lux,
                               pump_on, fan_on). None/[] = không có dữ liệu cảm biến.

    Returns
    -------
    bytes  – nội dung file .xlsx
    """
    if week_start is None:
        week_start = datetime.now() - timedelta(days=7)
    week_end = week_start + timedelta(days=6)
    sensor_records = sensor_records or []

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # xóa sheet mặc định

    _sheet_summary(wb, records, week_start, week_end)
    _sheet_sensor(wb, sensor_records)
    _sheet_daily(wb, records)
    _sheet_grid_heatmap(wb, records)
    _sheet_raw(wb, records)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_and_send_weekly_report(db_fetch_fn, week_start=None):
    """
    Wrapper gọi từ scheduler hoặc Flask route.

    Parameters
    ----------
    db_fetch_fn : callable  – hàm lấy records từ DB, nhận (start, end) -> list
    week_start  : datetime  – None = tuần hiện tại
    """
    if week_start is None:
        today = datetime.now()
        week_start = today - timedelta(days=today.weekday() + 1)  # Chủ nhật trước
        week_start = week_start.replace(hour=0, minute=0, second=0, microsecond=0)

    week_end = week_start + timedelta(days=6, hours=23, minutes=59)

    try:
        records = db_fetch_fn(week_start, week_end)
        xlsx_bytes = generate_weekly_excel(records, week_start)
        filename = f"BaoCao_CaiNgot_{week_start.strftime('%d%m%Y')}-{week_end.strftime('%d%m%Y')}.xlsx"
        _send_email(xlsx_bytes, filename, week_start, week_end, len(records))
        logger.info(f"Đã gửi báo cáo tuần: {filename}")
        return True, filename
    except Exception as e:
        logger.error(f"Lỗi tạo/gửi báo cáo: {e}")
        return False, str(e)


# ────────────────────────────────────────────────
# SHEET 1: TỔNG QUAN TUẦN
# ────────────────────────────────────────────────

def _sheet_summary(wb, records, week_start, week_end):
    ws = wb.create_sheet("📊 Tổng quan tuần")
    ws.sheet_view.showGridLines = False

    # --- Tiêu đề chính ---
    ws.merge_cells("A1:H1")
    ws["A1"] = "BÁO CÁO THEO DÕI SINH TRƯỞNG CẢI NGỌT"
    ws["A1"].font = _font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = _center()
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Tuần: {week_start.strftime('%d/%m/%Y')} – {week_end.strftime('%d/%m/%Y')}   |   Tổng ảnh phân tích: {len(records)}"
    ws["A2"].font = _font(size=10, color="4CAF50")
    ws["A2"].alignment = _center()
    ws.row_dimensions[2].height = 20

    # --- Thống kê tổng hợp ---
    if records:
        avg_plant   = round(sum(r.get("plant_count", 0) for r in records) / len(records), 1)
        avg_canopy  = round(sum(r.get("avg_canopy_cm2", 0) for r in records) / len(records), 2)
        total_warn  = sum(1 for r in records if r.get("disease_warning", False))
        last_health = records[-1].get("health_summary", {})
        health_score = last_health.get("health_score", 0)
    else:
        avg_plant = avg_canopy = total_warn = health_score = 0

    stats = [
        ("🌱 Số cây TB/lần chụp", f"{avg_plant} / 16"),
        ("🍃 Diện tích tán TB",   f"{avg_canopy} cm²"),
        ("⚠️  Lần cảnh báo bệnh", f"{total_warn} lần"),
        ("💚 Điểm sức khỏe cuối tuần", f"{health_score} %"),
    ]

    ws.row_dimensions[3].height = 10
    ws.merge_cells("A4:B4"); ws["A4"] = "CHỈ TIÊU"; ws["A4"].font = _font(bold=True, size=10, color="FFFFFF"); ws["A4"].fill = SUBHEAD_FILL
    ws.merge_cells("C4:D4"); ws["C4"] = "GIÁ TRỊ";  ws["C4"].font = _font(bold=True); ws["C4"].fill = SUBHEAD_FILL; ws["C4"].alignment = _center()

    for i, (label, val) in enumerate(stats, start=5):
        fill = ALT_ROW_FILL if i % 2 == 1 else _fill("FFFFFF")
        ws.merge_cells(f"A{i}:B{i}"); ws[f"A{i}"] = label; ws[f"A{i}"].fill = fill; ws[f"A{i}"].font = _font()
        ws.merge_cells(f"C{i}:D{i}"); ws[f"C{i}"] = val;   ws[f"C{i}"].fill = fill; ws[f"C{i}"].font = _font(bold=True); ws[f"C{i}"].alignment = _center()
        for col in ["A","B","C","D"]:
            ws[f"{col}{i}"].border = _border()
    ws.row_dimensions[3].height = 6

    # Cột rộng
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 18

    # --- Biểu đồ số cây theo ngày ---
    row_offset = 11
    ws[f"A{row_offset}"] = "Ngày"
    ws[f"B{row_offset}"] = "Số cây"
    ws[f"C{row_offset}"] = "Tán lá TB (cm²)"
    ws[f"A{row_offset}"].font = ws[f"B{row_offset}"].font = ws[f"C{row_offset}"].font = _font(bold=True)

    # Gom theo ngày
    daily = {}
    for r in records:
        ts  = r.get("timestamp", "")
        day = ts[:10] if ts else "N/A"
        if day not in daily:
            daily[day] = {"plants": [], "canopy": []}
        daily[day]["plants"].append(r.get("plant_count", 0))
        daily[day]["canopy"].append(r.get("avg_canopy_cm2", 0))

    chart_data_start = row_offset + 1
    for i, (day, vals) in enumerate(sorted(daily.items())):
        row = row_offset + 1 + i
        ws[f"A{row}"] = day
        ws[f"B{row}"] = round(sum(vals["plants"]) / len(vals["plants"]), 1)
        ws[f"C{row}"] = round(sum(vals["canopy"]) / len(vals["canopy"]), 2)
        ws[f"A{row}"].alignment = _center()

    chart_data_end = row_offset + len(daily)

    if len(daily) >= 2:
        chart = LineChart()
        chart.title = "Số cây & Tán lá trung bình theo ngày"
        chart.style = 10
        chart.height = 12
        chart.width  = 22

        data_plant = Reference(ws, min_col=2, min_row=row_offset, max_row=chart_data_end)
        data_can   = Reference(ws, min_col=3, min_row=row_offset, max_row=chart_data_end)
        cats       = Reference(ws, min_col=1, min_row=row_offset+1, max_row=chart_data_end)
        chart.add_data(data_plant, titles_from_data=True)
        chart.add_data(data_can,   titles_from_data=True)
        chart.set_categories(cats)
        chart.series[0].graphicalProperties.line.solidFill = "2E7D32"
        chart.series[1].graphicalProperties.line.solidFill = "FFA000"

        ws.add_chart(chart, f"E{row_offset}")


# ────────────────────────────────────────────────
# SHEET 2: CẢM BIẾN MÔI TRƯỜNG
# ────────────────────────────────────────────────

def _sheet_sensor(wb, sensor_records):
    ws = wb.create_sheet("🌡️ Cảm biến môi trường")
    ws.sheet_view.showGridLines = False

    headers = ["Thời gian", "Nhiệt độ (°C)", "Độ ẩm KK (%)",
               "Độ ẩm đất (%)", "Ánh sáng (lux)", "Bơm", "Quạt"]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = _border()
        cell.alignment = _center()
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    widths = [20, 14, 14, 14, 14, 10, 10]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    if not sensor_records:
        ws["A2"] = "(Không có dữ liệu cảm biến trong khoảng thời gian này)"
        ws["A2"].font = _font(size=10)
        return

    row = 2
    for r in sensor_records:
        ts = str(r.get("timestamp", ""))[:19].replace("T", " ")
        row_data = [
            ts,
            r.get("temperature"),
            r.get("humidity"),
            r.get("soil_moist"),
            r.get("light_lux"),
            "BẬT" if r.get("pump_on") else "TẮT",
            "BẬT" if r.get("fan_on") else "TẮT",
        ]
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=ci, value=val)
            cell.border = _border()
            cell.alignment = _center()
            if row % 2 == 0:
                cell.fill = ALT_ROW_FILL
        row += 1

    # Thống kê min/max/avg ở cuối sheet
    def _stats(key):
        vals = [r.get(key) for r in sensor_records if r.get(key) is not None]
        if not vals:
            return "—", "—", "—"
        return (round(min(vals), 1), round(max(vals), 1), round(sum(vals) / len(vals), 1))

    stat_row = row + 1
    ws.cell(row=stat_row, column=1, value="Min / Max / TB").font = _font(bold=True)
    for ci, key in [(2, "temperature"), (3, "humidity"), (4, "soil_moist"), (5, "light_lux")]:
        mn, mx, avg = _stats(key)
        cell = ws.cell(row=stat_row, column=ci, value=f"{mn} / {mx} / {avg}")
        cell.font = _font(bold=True, size=9)
        cell.fill = SUBHEAD_FILL
        cell.alignment = _center()


# ────────────────────────────────────────────────
# SHEET 3: CHI TIẾT TỪNG LẦN CHỤP
# ────────────────────────────────────────────────

def _sheet_daily(wb, records):
    ws = wb.create_sheet("📋 Chi tiết lần chụp")
    ws.sheet_view.showGridLines = False

    headers = [
        "STT", "Thời gian", "Số cây", "Thiếu vị trí",
        "Tán lá TB (cm²)", "Healthy", "Yellow", "Diseased",
        "Điểm SK (%)", "Cảnh báo bệnh"
    ]
    col_widths = [5, 20, 9, 28, 16, 10, 10, 10, 12, 14]

    # Header row
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font      = _font(bold=True, color="FFFFFF")
        cell.fill      = HEADER_FILL
        cell.alignment = _center()
        cell.border    = _border()
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    for ri, rec in enumerate(records, start=2):
        hs = rec.get("health_summary", {})
        missing_str = ", ".join(rec.get("missing_positions", [])) or "—"
        warn = rec.get("disease_warning", False)

        row_data = [
            ri - 1,
            rec.get("timestamp", "")[:19].replace("T", " "),
            rec.get("plant_count", 0),
            missing_str,
            rec.get("avg_canopy_cm2", 0),
            hs.get("healthy", 0),
            hs.get("yellow", 0),
            hs.get("diseased", 0),
            hs.get("health_score", 0),
            "⚠️ CÓ" if warn else "✓ Không",
        ]
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = _border()
            cell.alignment = _center()

            if ci == 10 and warn:
                cell.fill = DANGER_FILL
                cell.font = _font(bold=True, color="C62828")
            elif ri % 2 == 0:
                cell.fill = ALT_ROW_FILL

            # Tô màu cột điểm SK
            if ci == 9 and isinstance(val, (int, float)):
                if val >= 80:
                    cell.fill = HEALTHY_FILL
                elif val >= 60:
                    cell.fill = WARNING_FILL
                else:
                    cell.fill = DANGER_FILL


# ────────────────────────────────────────────────
# SHEET 4: HEATMAP LƯỚI 4×4
# ────────────────────────────────────────────────

def _sheet_grid_heatmap(wb, records):
    ws = wb.create_sheet("🗺️ Heatmap lưới cây")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:E1")
    ws["A1"] = "HEATMAP VỊ TRÍ CÂY – TRUNG BÌNH CẢ TUẦN"
    ws["A1"].font = _font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = _center()
    ws.row_dimensions[1].height = 28

    # Tính trung bình tán lá và sức khỏe theo vị trí
    pos_data = {}   # "R1C1" -> {"canopy": [], "health": []}
    for rec in records:
        for cr in rec.get("cell_results", []):
            key = cr.get("position_label", "")
            if key not in pos_data:
                pos_data[key] = {"canopy": [], "health_score": []}
            if cr.get("has_plant"):
                pos_data[key]["canopy"].append(cr.get("canopy_cm2", 0))
            # health score per cell: healthy=1, yellow=0.5, diseased=0, missing=0
            h_map = {"healthy": 1.0, "yellow": 0.5, "diseased": 0.0, "missing": 0.0}
            pos_data[key]["health_score"].append(h_map.get(cr.get("health", "missing"), 0))

    # --- Vẽ 2 bảng: Tán lá & Sức khỏe ---
    for table_idx, (title, metric) in enumerate([
        ("Diện tích tán lá TB (cm²)", "canopy"),
        ("Điểm sức khỏe TB (0–1)", "health_score"),
    ]):
        row_base = 3 + table_idx * 8

        ws.merge_cells(f"B{row_base}:E{row_base}")
        ws[f"B{row_base}"] = title
        ws[f"B{row_base}"].font = _font(bold=True, size=11)
        ws[f"B{row_base}"].fill = SUBHEAD_FILL
        ws[f"B{row_base}"].alignment = _center()

        # Cột headers (C1-C4)
        ws[f"A{row_base+1}"] = "↓ Hàng \\ Cột →"
        ws[f"A{row_base+1}"].font = _font(bold=True)
        for c in range(1, 5):
            cell = ws.cell(row=row_base+1, column=c+1, value=f"Cột {c}")
            cell.font = _font(bold=True); cell.fill = SUBHEAD_FILL; cell.alignment = _center()

        # Dữ liệu
        for r in range(1, 5):
            ws.cell(row=row_base+1+r, column=1, value=f"Hàng {r}").font = _font(bold=True)
            for c in range(1, 5):
                key = f"R{r}C{c}"
                vals = pos_data.get(key, {}).get(metric, [])
                avg  = round(sum(vals) / len(vals), 2) if vals else 0.0
                cell = ws.cell(row=row_base+1+r, column=c+1, value=avg)
                cell.alignment = _center()
                cell.border    = _border()

                # Gradient màu: xanh = tốt, đỏ = kém
                if metric == "health_score":
                    if avg >= 0.8:   cell.fill = HEALTHY_FILL
                    elif avg >= 0.5: cell.fill = WARNING_FILL
                    else:            cell.fill = DANGER_FILL
                else:  # canopy
                    all_vals = [v for d in pos_data.values() for v in d.get("canopy", [])]
                    if all_vals:
                        mn, mx = min(all_vals), max(all_vals)
                        ratio  = (avg - mn) / (mx - mn + 0.001)
                        g = int(100 + ratio * 120)
                        r_ch = int(220 - ratio * 160)
                        cell.fill = _fill(f"{r_ch:02X}{g:02X}50")

    # Chú thích
    ws["B18"] = "🟢 Xanh = tốt / nhiều   🟡 Vàng = trung bình   🔴 Đỏ = thiếu / bệnh"
    ws["B18"].font = _font(size=9)

    for col_idx in range(1, 6):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    for r in range(1, 20):
        ws.row_dimensions[r].height = 22


# ────────────────────────────────────────────────
# SHEET 5: RAW DATA
# ────────────────────────────────────────────────

def _sheet_raw(wb, records):
    ws = wb.create_sheet("📁 Dữ liệu thô")
    headers = [
        "Timestamp", "Vị trí", "Hàng", "Cột",
        "Có cây", "Sức khỏe", "Tán lá (cm²)",
        "Green ratio", "Yellow ratio", "Brown ratio", "Sick ratio"
    ]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = _border()
        ws.column_dimensions[get_column_letter(ci)].width = 16
    ws.freeze_panes = "A2"

    row = 2
    for rec in records:
        ts = rec.get("timestamp", "")[:19].replace("T", " ")
        for cr in rec.get("cell_results", []):
            row_data = [
                ts,
                cr.get("position_label", ""),
                cr.get("grid_row", ""),
                cr.get("grid_col", ""),
                "Có" if cr.get("has_plant") else "Không",
                cr.get("health", ""),
                cr.get("canopy_cm2", 0),
                cr.get("green_ratio", 0),
                cr.get("yellow_ratio", 0),
                cr.get("brown_ratio", 0),
                cr.get("sick_ratio", 0),
            ]
            for ci, val in enumerate(row_data, 1):
                cell = ws.cell(row=row, column=ci, value=val)
                cell.border = _border()
                if row % 2 == 0:
                    cell.fill = ALT_ROW_FILL
            row += 1


# ────────────────────────────────────────────────
# GỬI EMAIL
# ────────────────────────────────────────────────

def _send_email(xlsx_bytes: bytes, filename: str, week_start, week_end, record_count: int):
    msg = MIMEMultipart()
    msg["From"]    = EMAIL_SENDER
    msg["To"]      = EMAIL_RECIPIENT
    msg["Subject"] = f"[CẢI NGỌT IoT] Báo cáo tuần {week_start.strftime('%d/%m')}–{week_end.strftime('%d/%m/%Y')}"

    body = f"""Xin chào,

Hệ thống giám sát cải ngọt đã tổng hợp báo cáo tuần từ {week_start.strftime('%d/%m/%Y')} đến {week_end.strftime('%d/%m/%Y')}.

📊 Tổng số lần phân tích ảnh: {record_count}
📁 File đính kèm: {filename}

File Excel bao gồm:
  • Tổng quan tuần (biểu đồ số cây & tán lá theo ngày)
  • Cảm biến môi trường (nhiệt độ, độ ẩm không khí, độ ẩm đất, ánh sáng)
  • Chi tiết từng lần chụp ảnh
  • Heatmap lưới 4×4 vị trí cây
  • Dữ liệu thô để phân tích thêm

Trân trọng,
Hệ thống IoT Giám sát Cải Ngọt
"""
    msg.attach(MIMEText(body, "plain", "utf-8"))

    # Đính kèm file xlsx
    part = MIMEBase("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    part.set_payload(xlsx_bytes)
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f'attachment; filename="{filename}"')
    msg.attach(part)

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
        server.ehlo()
        server.starttls()
        server.login(EMAIL_SENDER, EMAIL_PASSWORD)
        server.sendmail(EMAIL_SENDER, EMAIL_RECIPIENT, msg.as_string())
